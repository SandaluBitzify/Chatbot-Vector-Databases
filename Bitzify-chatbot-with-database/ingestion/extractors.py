import fitz
import pandas as pd
from docx import Document
import re
from .pdf_extractors import extract_from_pdf_comprehensive

def extract_from_pdf(path):
    """Extract text from PDF using comprehensive extraction."""
    text, metadata = extract_from_pdf_comprehensive(path)
    return text

def extract_from_excel(path):
    """Extract text from Excel with better structure."""
    excel_file = pd.ExcelFile(path)
    all_text = []
    
    for sheet_name in excel_file.sheet_names:
        df = pd.read_excel(path, sheet_name=sheet_name)
        
        # Skip empty sheets
        if df.empty:
            continue
        
        sheet_text = f"\n=== Sheet: {sheet_name} ===\n"
        
        # Add column information
        columns = [str(col) for col in df.columns]
        sheet_text += f"Columns: {' | '.join(columns)}\n\n"
        
        # Process each row with better formatting
        for index, row in df.iterrows():
            row_data = []
            for col, val in zip(df.columns, row.values):
                if pd.notna(val):  # Only include non-null values
                    row_data.append(f"{col}: {str(val).strip()}")
            
            if row_data:  # Only add rows with data
                sheet_text += f"Record {index + 1}: {' | '.join(row_data)}\n"
        
        all_text.append(sheet_text)
    
    return "\n\n".join(all_text)

def extract_from_csv(path):
    """Extract text from CSV with better structure."""
    try:
        df = pd.read_csv(path)
        
        if df.empty:
            return "Empty CSV file"
        
        text = f"CSV Data from {path}\n"
        text += f"Total Records: {len(df)}\n"
        text += f"Columns: {' | '.join([str(col) for col in df.columns])}\n\n"
        
        # Process each row
        for index, row in df.iterrows():
            row_data = []
            for col, val in zip(df.columns, row.values):
                if pd.notna(val):
                    row_data.append(f"{col}: {str(val).strip()}")
            
            if row_data:
                text += f"Record {index + 1}: {' | '.join(row_data)}\n"
        
        return text
        
    except Exception as e:
        return f"Error reading CSV: {str(e)}"

def extract_from_docx(path):
    """Extract text from DOCX with better formatting."""
    doc = Document(path)
    text = ""
    
    # Extract paragraphs
    for para in doc.paragraphs:
        if para.text.strip():
            text += para.text.strip() + "\n"
    
    # Extract tables
    for table_num, table in enumerate(doc.tables):
        text += f"\n--- Table {table_num + 1} ---\n"
        
        # Get headers from first row
        if table.rows:
            headers = [cell.text.strip() for cell in table.rows[0].cells]
            text += f"Headers: {' | '.join(headers)}\n"
            
            # Process data rows
            for row_num, row in enumerate(table.rows[1:], 1):
                row_data = [cell.text.strip() for cell in row.cells]
                if any(data for data in row_data):  # Only add non-empty rows
                    formatted_row = ' | '.join(f"{h}: {d}" for h, d in zip(headers, row_data) if d)
                    text += f"Row {row_num}: {formatted_row}\n"
    
    return clean_extracted_text(text)

def extract_from_txt(path):
    """Extract text from TXT file."""
    try:
        with open(path, 'r', encoding='utf-8') as f:
            content = f.read()
        return clean_extracted_text(content)
    except UnicodeDecodeError:
        # Try with different encoding
        with open(path, 'r', encoding='latin-1') as f:
            content = f.read()
        return clean_extracted_text(content)

def clean_extracted_text(text):
    """Clean and normalize extracted text."""
    if not text:
        return ""
    
    # Remove excessive whitespace
    text = re.sub(r'\s+', ' ', text)
    
    # Remove special characters that cause issues
    text = re.sub(r'[^\w\s.,!?;:()\-\[\]{}"/|$%]', ' ', text)
    
    # Fix common OCR issues
    text = text.replace(' . ', '. ')
    text = text.replace(' , ', ', ')
    
    return text.strip()

def extract_from_excel_with_metadata(path):
    """Extract text from Excel with complete metadata including record counts."""
    print(f"📊 Starting Excel extraction from: {path}")
    
    try:
        excel_file = pd.ExcelFile(path)
        print(f"📊 Found {len(excel_file.sheet_names)} sheets: {excel_file.sheet_names}")
        
        all_text = []
        total_records = 0
        all_columns = []
        sheet_info = []
        
        for sheet_name in excel_file.sheet_names:
            try:
                print(f"📊 Processing sheet: {sheet_name}")
                df = pd.read_excel(path, sheet_name=sheet_name)
                
                # Skip empty sheets
                if df.empty:
                    print(f"⚠️ Sheet {sheet_name} is empty, skipping")
                    continue
                
                sheet_records = len(df)
                total_records += sheet_records
                print(f"📊 Sheet {sheet_name} has {sheet_records} records")
                
                # Store sheet information
                sheet_info.append({
                    'name': sheet_name,
                    'records': sheet_records,
                    'columns': [str(col) for col in df.columns]
                })
                
                # Collect all unique columns
                sheet_columns = [str(col) for col in df.columns]
                all_columns.extend(sheet_columns)
                print(f"📊 Sheet columns: {sheet_columns}")
                
                # Create text representation
                sheet_text = f"\n=== Sheet: {sheet_name} ===\n"
                sheet_text += f"Total Records in Sheet: {sheet_records}\n"
                sheet_text += f"Columns: {' | '.join(sheet_columns)}\n\n"
                
                # Add sample data (first 10 and last 5 rows to give context)
                sample_size = min(10, len(df))
                for index, row in df.head(sample_size).iterrows():
                    row_data = []
                    for col, val in zip(df.columns, row.values):
                        if pd.notna(val):
                            row_data.append(f"{col}: {str(val).strip()}")
                    
                    if row_data:
                        sheet_text += f"Record {index + 1}: {' | '.join(row_data)}\n"
                
                # If there are many records, add some from the end too
                if len(df) > 15:
                    sheet_text += f"\n... [Skipping {len(df) - 15} records] ...\n\n"
                    for index, row in df.tail(5).iterrows():
                        row_data = []
                        for col, val in zip(df.columns, row.values):
                            if pd.notna(val):
                                row_data.append(f"{col}: {str(val).strip()}")
                        
                        if row_data:
                            sheet_text += f"Record {index + 1}: {' | '.join(row_data)}\n"
                
                all_text.append(sheet_text)
                
            except Exception as e:
                print(f"❌ Error processing sheet {sheet_name}: {e}")
                continue
        
        # Create metadata
        metadata = {
            'total_records': total_records,
            'columns': list(set(all_columns)),
            'sheets': sheet_info,
            'total_sheets': len(sheet_info)
        }
        
        print(f"📊 Final Excel metadata: {metadata}")
        
        combined_text = "\n\n".join(all_text)
        
        # Add summary at the beginning
        summary = f"EXCEL FILE SUMMARY:\n"
        summary += f"Total Records Across All Sheets: {total_records}\n"
        summary += f"Total Sheets: {len(sheet_info)}\n"
        summary += f"Sheet Details: {', '.join([f'{s['name']} ({s['records']} records)' for s in sheet_info])}\n\n"


        
        final_text = summary + combined_text
        print(f"📊 Generated text length: {len(final_text)} characters")
        
        return final_text, metadata
        
    except Exception as e:
        print(f"❌ Error in extract_from_excel_with_metadata: {e}")
        import traceback
        traceback.print_exc()
        raise

def extract_from_csv_with_metadata(path):
    """Extract text from CSV with metadata."""
    print(f"📊 Starting CSV extraction from: {path}")
    
    try:
        df = pd.read_csv(path)
        
        if df.empty:
            return "Empty CSV file", {'total_records': 0, 'columns': []}
        
        total_records = len(df)
        columns = [str(col) for col in df.columns]
        
        print(f"📊 CSV has {total_records} records and {len(columns)} columns")
        
        # Create summary
        text = f"CSV FILE SUMMARY:\n"
        text += f"Total Records: {total_records}\n"
        text += f"Columns ({len(columns)}): {' | '.join(columns)}\n\n"
        
        # Add sample data
        sample_size = min(10, len(df))
        for index, row in df.head(sample_size).iterrows():
            row_data = []
            for col, val in zip(df.columns, row.values):
                if pd.notna(val):
                    row_data.append(f"{col}: {str(val).strip()}")
            
            if row_data:
                text += f"Record {index + 1}: {' | '.join(row_data)}\n"
        
        if len(df) > 10:
            text += f"\n... [Total of {total_records} records in file] ...\n"
        
        metadata = {
            'total_records': total_records,
            'columns': columns
        }
        
        return text, metadata
        
    except Exception as e:
        print(f"❌ Error in extract_from_csv_with_metadata: {e}")
        return f"Error reading CSV: {str(e)}", {'total_records': 0, 'columns': []}

def extract_from_pdf_with_metadata(path):
    """Extract text from PDF with comprehensive metadata."""
    print(f"📄 Starting comprehensive PDF extraction from: {path}")
    
    try:
        text, metadata = extract_from_pdf_comprehensive(path)
        
        # Add file-level metadata
        metadata.update({
            'file_type': 'pdf',
            'columns': [],  # PDFs don't have columns like spreadsheets
            'extraction_summary': f"Used {len(metadata.get('extraction_methods', []))} extraction methods"
        })
        
        return text, metadata
        
    except Exception as e:
        print(f"❌ Error in PDF extraction: {e}")
        import traceback
        traceback.print_exc()
        raise
